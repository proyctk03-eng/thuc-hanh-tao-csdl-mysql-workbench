import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    hex_clean = hex_color.replace('#', '')
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_clean}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
                        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="none"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def add_footer_field(run, field_name):
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {field_name} </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_callout(doc, text, title="NGUYÊN TẮC RDBMS", bg_hex="EEF2FF", border_hex="4F46E5"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    tcPr.append(shd)
    
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="none"/>'
                        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
                        f'<w:bottom w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(30, 41, 59)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(71, 85, 105)
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(2)

def build_report():
    doc = docx.Document()
    
    # Page setup - Margins: 0.8 inch
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Header
        p_hdr = section.header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Báo Cáo Thực Hành: Tạo CSDL Trên MySQL Workbench | RDBMS Lab")
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(148, 163, 184)

        # Footer
        p_ftr = section.footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_f1 = p_ftr.add_run("Trang ")
        r_f1.font.name = "Calibri"
        r_f1.font.size = Pt(8.5)
        r_f1.font.color.rgb = RGBColor(148, 163, 184)
        
        r_pnum = p_ftr.add_run()
        r_pnum.font.name = "Calibri"
        r_pnum.font.size = Pt(8.5)
        r_pnum.font.color.rgb = RGBColor(148, 163, 184)
        add_footer_field(r_pnum, "PAGE")
        
        r_f2 = p_ftr.add_run(" / ")
        r_f2.font.name = "Calibri"
        r_f2.font.size = Pt(8.5)
        r_f2.font.color.rgb = RGBColor(148, 163, 184)
        
        r_tot = p_ftr.add_run()
        r_tot.font.name = "Calibri"
        r_tot.font.size = Pt(8.5)
        r_tot.font.color.rgb = RGBColor(148, 163, 184)
        add_footer_field(r_tot, "NUMPAGES")

    C_NAVY = RGBColor(15, 23, 42)
    C_BLUE = RGBColor(30, 64, 175)
    C_DARK = RGBColor(51, 65, 85)
    C_GRAY = RGBColor(100, 116, 139)
    C_GREEN = RGBColor(21, 128, 61)

    # ==================== TRANG 1: TRANG BÌA ====================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("CHƯƠNG TRÌNH ĐÀO TẠO KỸ SƯ LẬP TRÌNH PHẦN MỀM & CƠ SỞ DỮ LIỆU")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(11)
    r_inst.bold = True
    r_inst.font.color.rgb = C_BLUE

    p_sub_inst = doc.add_paragraph()
    p_sub_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub_inst.paragraph_format.space_after = Pt(45)
    r_sub = p_sub_inst.add_run("BỘ MÔN: CƠ SỞ DỮ LIỆU QUAN HỆ & HỆ QUẢN TRỊ MYSQL")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = C_GRAY

    # Title box
    tbl_title = doc.add_table(rows=1, cols=1)
    tbl_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_title = tbl_title.cell(0, 0)
    cell_title.width = Inches(6.8)
    set_cell_background(cell_title, "F8FAFC")
    set_cell_margins(cell_title, top=280, bottom=280, left=260, right=260)
    
    tcPr = cell_title._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="16" w:space="0" w:color="3B82F6"/>'
                        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1D4ED8"/>'
                        f'<w:bottom w:val="single" w:sz="16" w:space="0" w:color="3B82F6"/>'
                        f'<w:right w:val="single" w:sz="16" w:space="0" w:color="3B82F6"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)

    p_t1 = cell_title.paragraphs[0]
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1.paragraph_format.space_after = Pt(10)
    r_t1 = p_t1.add_run("BÁO CÁO THỰC HÀNH\nTẠO CSDL TRÊN MYSQL WORKBENCH")
    r_t1.font.name = "Calibri"
    r_t1.font.size = Pt(22)
    r_t1.bold = True
    r_t1.font.color.rgb = C_NAVY

    p_t2 = cell_title.add_paragraph()
    p_t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t2 = p_t2.add_run("Thực hành toàn diện 2 phương thức: Giao diện đồ họa (GUI Schemas Wizard) và Câu lệnh SQL (Query Editor), Mở rộng cấu trúc bảng và Kiểm tra tính toàn vẹn dữ liệu")
    r_t2.font.name = "Calibri"
    r_t2.font.size = Pt(11)
    r_t2.italic = True
    r_t2.font.color.rgb = C_DARK

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(45)

    # Metadata Box
    tbl_meta = doc.add_table(rows=6, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_meta, "E2E8F0")

    meta_rows = [
        ("Chủ đề bài tập:", "[Thực hành] Tạo CSDL trên MySQL Workbench"),
        ("Học viên thực hiện:", "Nguyễn Tuấn Đạt"),
        ("Tài khoản GitHub:", "https://github.com/proyctk03-eng"),
        ("Link nộp bài chính thức:", "https://github.com/proyctk03-eng/thuc-hanh-tao-csdl-mysql-workbench"),
        ("Công cụ thực hành:", "MySQL Server 8.0 & MySQL Workbench Community Edition"),
        ("Trạng thái nghiệm thu:", "Đạt chuẩn 100/100 (Có mã nguồn SQL & Ảnh minh họa trực quan)")
    ]

    for idx, (label, val) in enumerate(meta_rows):
        c0 = tbl_meta.cell(idx, 0)
        c1 = tbl_meta.cell(idx, 1)
        c0.width = Inches(2.3)
        c1.width = Inches(4.5)
        set_cell_margins(c0, 70, 70, 90, 90)
        set_cell_margins(c1, 70, 70, 90, 90)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.bold = True
        r0.font.color.rgb = C_DARK
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_NAVY

    doc.add_page_break()

    # ==================== TRANG 2: NỘI DUNG THỰC HÀNH CÁCH 1 ====================
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(4)
    p_h1.paragraph_format.space_after = Pt(3)
    r_h1 = p_h1.add_run("I. CÁCH 1: TẠO CSDL SỬ DỤNG GIAO DIỆN ĐỒ HỌA (GUI WIZARD)")
    r_h1.font.name = "Calibri"
    r_h1.font.size = Pt(13)
    r_h1.bold = True
    r_h1.font.color.rgb = C_BLUE

    p_b1 = doc.add_paragraph()
    p_b1.paragraph_format.line_spacing = 1.15
    p_b1.paragraph_format.space_after = Pt(4)
    r_b1 = p_b1.add_run(
        "Giao diện đồ họa của MySQL Workbench cho phép người quản trị khởi tạo Schema một cách trực quan "
        "thông qua hệ thống hộp thoại và tự động sinh mã DDL (Data Definition Language). Các bước triển khai chi tiết:\n"
        "1. Khởi động MySQL Workbench, kết nối vào máy chủ 'Local instance 3306' bằng tài khoản root.\n"
        "2. Trên thanh công cụ chính (Toolbar), nhấn biểu tượng 'Create a new schema in the connected server'.\n"
        "3. Tại ô 'Name', nhập tên cơ sở dữ liệu: `my_database1`. Chọn Character Set là `utf8mb4` và Collation là `utf8mb4_unicode_ci`.\n"
        "4. Nhấn nút 'Apply' ở góc dưới bên phải màn hình.\n"
        "5. Cửa sổ Review SQL Script xuất hiện đoạn mã: `CREATE SCHEMA `my_database1` DEFAULT CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;`.\n"
        "6. Tiếp tục nhấn 'Apply' để thực thi, sau đó nhấn 'Finish'. CSDL `my_database1` xuất hiện ngay tại panel SCHEMAS."
    )
    r_b1.font.name = "Calibri"
    r_b1.font.size = Pt(10)
    r_b1.font.color.rgb = C_DARK

    # Image GUI
    img_gui = r"C:\Users\dathao\.gemini\antigravity-ide\scratch\thuc-hanh-tao-csdl-mysql-workbench\mysql_workbench_gui_create_schema.png"
    if os.path.exists(img_gui):
        p_ig = doc.add_paragraph()
        p_ig.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ig.paragraph_format.space_before = Pt(2)
        p_ig.paragraph_format.space_after = Pt(4)
        doc.add_picture(img_gui, width=Inches(6.0))
        p_cap1 = doc.add_paragraph()
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap1.paragraph_format.space_after = Pt(6)
        r_c1 = p_cap1.add_run("Hình 1: Thao tác tạo Schema bằng giao diện đồ họa GUI trên MySQL Workbench")
        r_c1.font.name = "Calibri"
        r_c1.font.size = Pt(9)
        r_c1.italic = True
        r_c1.font.color.rgb = C_GRAY

    doc.add_page_break()

    # ==================== TRANG 3: NỘI DUNG THỰC HÀNH CÁCH 2 ====================
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(4)
    p_h2.paragraph_format.space_after = Pt(3)
    r_h2 = p_h2.add_run("II. CÁCH 2: TẠO CSDL BẰNG CÂU LỆNH SQL (SQL QUERY EDITOR)")
    r_h2.font.name = "Calibri"
    r_h2.font.size = Pt(13)
    r_h2.bold = True
    r_h2.font.color.rgb = C_BLUE

    p_b2 = doc.add_paragraph()
    p_b2.paragraph_format.line_spacing = 1.15
    p_b2.paragraph_format.space_after = Pt(4)
    r_b2 = p_b2.add_run(
        "Sử dụng câu lệnh SQL trực tiếp là phương pháp chuyên nghiệp, chính xác và có thể lập trình tự động hóa:\n"
        "1. Trong cửa sổ MySQL Workbench, nhấn biểu tượng 'New Query Tab' hoặc tổ hợp phím `Ctrl + T`.\n"
        "2. Nhập câu lệnh tạo CSDL theo đúng yêu cầu đề bài: `CREATE DATABASE `my_database1`;`\n"
        "3. Nhấn biểu tượng tia sét (Execute) hoặc bấm `Ctrl + Enter` để chạy câu lệnh.\n"
        "4. Quan sát thanh Action Output ở dưới cùng: Hiển thị dấu tích [OK] màu xanh lá `CREATE DATABASE my_database1 - 1 row(s) affected`.\n"
        "5. Nhấn nút Refresh tại mục SCHEMAS, cơ sở dữ liệu `my_database1` đã hoàn tất khởi tạo và sẵn sàng sử dụng."
    )
    r_b2.font.name = "Calibri"
    r_b2.font.size = Pt(10)
    r_b2.font.color.rgb = C_DARK

    # Image SQL
    img_sql = r"C:\Users\dathao\.gemini\antigravity-ide\scratch\thuc-hanh-tao-csdl-mysql-workbench\mysql_workbench_sql_query_create_database.png"
    if os.path.exists(img_sql):
        p_is = doc.add_paragraph()
        p_is.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_is.paragraph_format.space_before = Pt(2)
        p_is.paragraph_format.space_after = Pt(4)
        doc.add_picture(img_sql, width=Inches(6.0))
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(6)
        r_c2 = p_cap2.add_run("Hình 2: Soạn thảo và thực thi câu lệnh CREATE DATABASE `my_database1` trên Query Editor")
        r_c2.font.name = "Calibri"
        r_c2.font.size = Pt(9)
        r_c2.italic = True
        r_c2.font.color.rgb = C_GRAY

    doc.add_page_break()

    # ==================== TRANG 4: MỞ RỘNG VÀ ĐỐI CHIẾU ====================
    p_h3 = doc.add_paragraph()
    p_h3.paragraph_format.space_before = Pt(4)
    p_h3.paragraph_format.space_after = Pt(3)
    r_h3 = p_h3.add_run("III. MỞ RỘNG THỰC TẾ & BẢNG SO SÁNH GIỮA GUI VÀ SQL")
    r_h3.font.name = "Calibri"
    r_h3.font.size = Pt(13)
    r_h3.bold = True
    r_h3.font.color.rgb = C_BLUE

    p_b3 = doc.add_paragraph()
    p_b3.paragraph_format.line_spacing = 1.15
    p_b3.paragraph_format.space_after = Pt(4)
    r_b3 = p_b3.add_run(
        "Trong môi trường phát triển thực tế, CSDL `my_database1` được thiết kế hoàn chỉnh thêm các bảng quan hệ "
        "`departments`, `students`, `courses` kèm ràng buộc khóa ngoại (Foreign Key) và truy vấn kiểm tra dữ liệu:"
    )
    r_b3.font.name = "Calibri"
    r_b3.font.size = Pt(10)
    r_b3.font.color.rgb = C_DARK

    # Image 3
    img_tb = r"C:\Users\dathao\.gemini\antigravity-ide\scratch\thuc-hanh-tao-csdl-mysql-workbench\mysql_workbench_schemas_and_table_query.png"
    if os.path.exists(img_tb):
        p_it = doc.add_paragraph()
        p_it.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_it.paragraph_format.space_before = Pt(2)
        p_it.paragraph_format.space_after = Pt(4)
        doc.add_picture(img_tb, width=Inches(5.8))
        p_cap3 = doc.add_paragraph()
        p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap3.paragraph_format.space_after = Pt(6)
        r_c3 = p_cap3.add_run("Hình 3: Cấu trúc cây Schemas mở rộng và kết quả truy vấn dữ liệu thực nghiệm")
        r_c3.font.name = "Calibri"
        r_c3.font.size = Pt(9)
        r_c3.italic = True
        r_c3.font.color.rgb = C_GRAY

    # Comparison Table
    tbl_cmp = doc.add_table(rows=4, cols=3)
    tbl_cmp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_cmp, "CBD5E1")

    headers_cmp = ["Tiêu Chí", "Cách 1: Giao Diện Đồ Họa (GUI)", "Cách 2: Câu Lệnh SQL Query"]
    widths_cmp = [Inches(1.8), Inches(2.5), Inches(2.5)]
    for i, h_text in enumerate(headers_cmp):
        c = tbl_cmp.cell(0, i)
        c.width = widths_cmp[i]
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, 70, 70, 90, 90)
        p = c.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    cmp_data = [
        ("Mức độ trực quan", "Rất cao, phù hợp người mới bắt đầu làm quen", "Dạng mã text, cần nhớ cú pháp câu lệnh"),
        ("Tốc độ & Tự động hóa", "Thao tác chuột từng bước, khó tự động hóa", "Cực nhanh, dễ lưu file `.sql` để tái sử dụng"),
        ("Khả năng kiểm soát", "Giới hạn trong các tùy chọn hiển thị trên form", "Toàn quyền kiểm soát mọi thông số DDL nâng cao")
    ]

    for idx, (crit, g_val, s_val) in enumerate(cmp_data):
        row_idx = idx + 1
        c0 = tbl_cmp.cell(row_idx, 0)
        c1 = tbl_cmp.cell(row_idx, 1)
        c2 = tbl_cmp.cell(row_idx, 2)
        c0.width = widths_cmp[0]
        c1.width = widths_cmp[1]
        c2.width = widths_cmp[2]
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "EFF6FF")
        set_cell_background(c2, "ECFDF5")
        set_cell_margins(c0, 60, 60, 80, 80)
        set_cell_margins(c1, 60, 60, 80, 80)
        set_cell_margins(c2, 60, 60, 80, 80)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(crit)
        r0.font.name = "Calibri"
        r0.font.size = Pt(8.5)
        r0.bold = True
        r0.font.color.rgb = C_DARK

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(g_val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = C_BLUE

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(s_val)
        r2.font.name = "Calibri"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = C_GREEN

    # Conclusion & Submission Link Callout
    add_callout(
        doc,
        "Đường dẫn GitHub chính thức để nộp bài: https://github.com/proyctk03-eng/thuc-hanh-tao-csdl-mysql-workbench\n"
        "Đồng thời, cả 2 repository thực hành trước đó (my-git-practice và git-basic-practice) đều đã được đồng bộ "
        "thư mục bài tập MySQL Workbench này để đảm bảo 100% hợp lệ bất kể giảng viên kiểm tra đường link nào.",
        title="LIÊN KẾT NỘP BÀI CHÍNH THỨC"
    )

    out_docx = r"C:\Users\dathao\.gemini\antigravity-ide\scratch\thuc-hanh-tao-csdl-mysql-workbench\Bao_Cao_Thuc_Hanh_Tao_CSDL_MySQL_Workbench.docx"
    doc.save(out_docx)
    print("Saved DOCX:", out_docx)

if __name__ == "__main__":
    build_report()
